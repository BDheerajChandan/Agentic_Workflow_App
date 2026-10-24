# app/files/readers.py
# Format-specific extraction. Every reader returns (content, results, extra_metadata).
# Heavy libraries are imported lazily so a missing optional dependency never breaks startup.

import csv
import io
import json
import shutil
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path

from app.files.registry import SOFFICE_TIMEOUT, FileError


def decode_text(data: bytes) -> str:
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8", errors="replace")
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        pass
    try:
        import chardet
        guess = chardet.detect(data[:200000])
        enc = guess.get("encoding")
        if enc and (guess.get("confidence") or 0) >= 0.5:
            return data.decode(enc, errors="replace")
    except ImportError:
        pass
    return data.decode("cp1252", errors="replace")


def _need(module_error: ImportError, package: str):
    raise FileError(f"This format requires the '{package}' package on the server.") from module_error


# ─── PDF ─────────────────────────────────────────────────────────────────────
def _pdf(path):
    try:
        import fitz
    except ImportError as e:
        _need(e, "pymupdf")
    pages = []
    with fitz.open(str(path), filetype="pdf") as doc:
        if doc.needs_pass:
            raise FileError("The PDF is password protected.")
        for i, page in enumerate(doc, 1):
            pages.append({"page_number": i, "content": page.get_text("text").strip()})
    text = "\n\n".join(p["content"] for p in pages if p["content"])
    meta = {"page_count": len(pages)}
    if pages and not text:
        meta["warning"] = "No extractable text found. The PDF may be scanned; OCR is not enabled."
    return text, {"pages": pages}, meta


# ─── DOCX ────────────────────────────────────────────────────────────────────
def _docx(path):
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as e:
        _need(e, "python-docx")
    doc = Document(str(path))
    blocks, paragraphs, tables = [], [], []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = Paragraph(child, doc).text.strip()
            if t:
                paragraphs.append(t)
                blocks.append(t)
        elif tag == "tbl":
            rows = [[c.text.strip() for c in r.cells] for r in Table(child, doc).rows]
            tables.append({"table_number": len(tables) + 1, "rows": rows})
            blocks.append("\n".join(" | ".join(r) for r in rows))
    return (
        "\n\n".join(blocks),
        {"paragraphs": paragraphs, "tables": tables},
        {"paragraph_count": len(paragraphs), "table_count": len(tables)},
    )


# ─── PPTX ────────────────────────────────────────────────────────────────────
def _walk_shapes(shapes, texts, tables):
    for shape in shapes:
        if hasattr(shape, "shapes"):
            _walk_shapes(shape.shapes, texts, tables)
        if getattr(shape, "has_table", False) and shape.has_table:
            rows = [[c.text.strip() for c in r.cells] for r in shape.table.rows]
            tables.append({"table_number": len(tables) + 1, "rows": rows})
            texts.append("\n".join(" | ".join(r) for r in rows))
        elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if t:
                texts.append(t)


def _pptx(path):
    try:
        from pptx import Presentation
    except ImportError as e:
        _need(e, "python-pptx")
    prs = Presentation(str(path))
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        texts, tables = [], []
        _walk_shapes(slide.shapes, texts, tables)
        slides.append({"slide_number": i, "content": "\n".join(texts), "tables": tables})
    text = "\n\n".join(f"[Slide {s['slide_number']}]\n{s['content']}" for s in slides if s["content"])
    return text, {"slides": slides}, {"slide_count": len(slides)}


# ─── Excel ───────────────────────────────────────────────────────────────────
def _excel(path, engine):
    try:
        import pandas as pd
        sheets = pd.read_excel(str(path), sheet_name=None, header=None, dtype=str, engine=engine)
    except ImportError as e:
        _need(e, "pandas / openpyxl / xlrd")
    out, parts = [], []
    for name, df in sheets.items():
        rows = df.fillna("").values.tolist()
        out.append({
            "sheet_name": str(name),
            "row_count": len(rows),
            "column_count": int(df.shape[1]),
            "rows": rows,
        })
        parts.append(f"[Sheet: {name}]\n" + "\n".join("\t".join(str(c) for c in r) for r in rows))
    return "\n\n".join(parts), {"sheets": out}, {"sheet_count": len(out)}


# ─── Simple text-based formats ───────────────────────────────────────────────
def _delimited(data, delimiter):
    csv.field_size_limit(min(sys.maxsize, 2**31 - 1))
    rows = list(csv.reader(io.StringIO(decode_text(data)), delimiter=delimiter))
    text = "\n".join(delimiter.join(r) for r in rows)
    return text, {"rows": rows}, {
        "row_count": len(rows),
        "column_count": max((len(r) for r in rows), default=0),
    }


def _rtf(data):
    try:
        from striprtf.striprtf import rtf_to_text
    except ImportError as e:
        _need(e, "striprtf")
    return rtf_to_text(decode_text(data)).strip(), {}, {}


def _html(data):
    try:
        from bs4 import BeautifulSoup
    except ImportError as e:
        _need(e, "beautifulsoup4")
    soup = BeautifulSoup(decode_text(data), "html.parser")
    title = soup.title.get_text(strip=True) if soup.title else None
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    return soup.get_text("\n", strip=True), {"title": title}, {}


def _xml(data):
    upper = data[:4096].upper()
    if b"<!DOCTYPE" in upper or b"<!ENTITY" in upper:
        raise FileError("XML with DOCTYPE/ENTITY declarations is not allowed.")
    try:
        root = ET.fromstring(data)
    except ET.ParseError as exc:
        raise FileError(f"Malformed XML: {exc}")
    text = "\n".join(t.strip() for t in root.itertext() if t and t.strip())
    return text, {"root_tag": root.tag}, {}


def _json(data):
    try:
        obj = json.loads(decode_text(data))
    except ValueError as exc:
        raise FileError(f"Malformed JSON: {exc}")
    return json.dumps(obj, ensure_ascii=False, indent=2), {"data": obj}, {}


def _jsonl(data):
    records = []
    for n, line in enumerate(decode_text(data).splitlines(), 1):
        if not line.strip():
            continue
        try:
            records.append(json.loads(line))
        except ValueError as exc:
            raise FileError(f"Malformed JSONL at line {n}: {exc}")
    text = "\n".join(json.dumps(r, ensure_ascii=False) for r in records)
    return text, {"records": records}, {"record_count": len(records)}


# ─── Legacy Office via LibreOffice ───────────────────────────────────────────
_LEGACY = {"doc": ("docx", _docx), "ppt": ("pptx", _pptx), "xls": ("xlsx", None)}


def _convert_legacy(path, kind):
    target, handler = _LEGACY[kind]
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        raise FileError("Legacy Office files need LibreOffice ('soffice') installed and on PATH.")
    with tempfile.TemporaryDirectory(prefix="fr_") as tmp:
        src = Path(tmp) / f"input.{kind}"
        shutil.copyfile(path, src)
        try:
            subprocess.run(
                [exe, "--headless", f"-env:UserInstallation={(Path(tmp) / 'profile').as_uri()}",
                 "--convert-to", target, "--outdir", tmp, str(src)],
                capture_output=True, timeout=SOFFICE_TIMEOUT, check=True,
            )
        except subprocess.TimeoutExpired:
            raise FileError("Legacy file conversion timed out.")
        except subprocess.CalledProcessError:
            raise FileError("Legacy file conversion failed.")
        out = Path(tmp) / f"input.{target}"
        if not out.exists():
            raise FileError("Legacy file conversion produced no output.")
        return handler(out) if handler else _excel(out, "openpyxl")


# ─── Entry point ─────────────────────────────────────────────────────────────
def extract(path, kind: str) -> dict:
    path = Path(path)
    if kind == "pdf":
        content, results, meta = _pdf(path)
    elif kind == "docx":
        content, results, meta = _docx(path)
    elif kind == "pptx":
        content, results, meta = _pptx(path)
    elif kind == "xlsx":
        content, results, meta = _excel(path, "openpyxl")
    elif kind == "xls":
        try:
            content, results, meta = _excel(path, "xlrd")
        except FileError:
            content, results, meta = _convert_legacy(path, "xls")
        except Exception:  # noqa: BLE001
            content, results, meta = _convert_legacy(path, "xls")
    elif kind in ("doc", "ppt"):
        content, results, meta = _convert_legacy(path, kind)
    else:
        data = path.read_bytes()
        if kind == "csv":
            content, results, meta = _delimited(data, ",")
        elif kind == "tsv":
            content, results, meta = _delimited(data, "\t")
        elif kind == "rtf":
            content, results, meta = _rtf(data)
        elif kind == "html":
            content, results, meta = _html(data)
        elif kind == "xml":
            content, results, meta = _xml(data)
        elif kind == "json":
            content, results, meta = _json(data)
        elif kind == "jsonl":
            content, results, meta = _jsonl(data)
        elif kind == "text":
            content, results, meta = decode_text(data), {}, {}
        else:
            raise FileError(f"No reader available for '{kind}'.")
    return {"content": content, "results": results, "meta": meta}